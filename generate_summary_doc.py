import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_summary_document():
    doc = docx.Document()

    # 1. Page Geometry (1-inch standard margins)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # 2. Color Palette
    PRIMARY = RGBColor(16, 44, 87)       # Navy
    SECONDARY = RGBColor(53, 89, 142)    # Slate Blue
    BODY = RGBColor(30, 30, 30)          # Off-black
    MUTED = RGBColor(110, 110, 110)      # Gray
    BORDER_HEX = "D3D3D3"
    HDR_BG_HEX = "102C57"

    # Helper: Paragraph Styling
    def add_p(text, style_type='Normal', space_after=6, bold=False, italic=False, color=BODY, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Calibri'
        run.font.size = Pt(size)
        run.font.color.rgb = color
        return p

    def add_h(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.bold = True
        if level == 1:
            run.font.size = Pt(15)
            run.font.color.rgb = PRIMARY
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
        elif level == 2:
            run.font.size = Pt(12.5)
            run.font.color.rgb = SECONDARY
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
        return p

    def add_code_block(code_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.left_indent = Inches(0.25)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(40, 40, 40)

    def style_cell(cell, bg_hex=None, text_color="000000", bold=False, font_size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
        tcPr = cell._tc.get_or_add_tcPr()
        if bg_hex:
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
            tcPr.append(shd)
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        for r in p.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(font_size)
            r.bold = bold
            r.font.color.rgb = RGBColor.from_string(text_color)

    def set_table_borders(table):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{BORDER_HEX}"/>'
            f'<w:left w:val="none"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{PRIMARY}"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{BORDER_HEX}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def create_table(headers, data, col_widths=None):
        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)

        # Header Row
        hdr_cells = table.rows[0].cells
        for i, title in enumerate(headers):
            hdr_cells[i].text = title
            style_cell(hdr_cells[i], bg_hex=HDR_BG_HEX, text_color="FFFFFF", bold=True, font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)

        # Data Rows
        for row_idx, row_values in enumerate(data):
            row_cells = table.add_row().cells
            bg = "F9FBFD" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, val in enumerate(row_values):
                row_cells[col_idx].text = str(val)
                align = WD_ALIGN_PARAGRAPH.CENTER if col_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
                style_cell(row_cells[col_idx], bg_hex=bg, text_color="202020", bold=False, font_size=9, align=align)

        if col_widths:
            for row in table.rows:
                for idx, width in enumerate(col_widths):
                    row.cells[idx].width = Inches(width)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==========================
    # DOCUMENT CONTENT
    # ==========================

    # Title & Header
    add_p("PROJECT LIFECYCLE SUMMARY", bold=True, size=22, color=PRIMARY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_p("120 Years of Olympic History: End-to-End SQL Analytics & Engineering", bold=True, size=13, color=SECONDARY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_p("Repository: github.com/dvloperhol8701/SportsStats-Olympic-SQL-Analytics | Complete Process & Technical Audit", italic=True, size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)

    # 1. Executive Overview
    add_h("1. Executive Overview & Problem Definition", level=1)
    add_p("The objective of the SportsStats Capstone project was to design, execute, and deliver an end-to-end data analytics pipeline evaluating 120 years of modern Olympic competition records (1896–2016). The project satisfied a dual-stakeholder engagement model: (1) High-Performance Training Academies demanding physiological archetypes, peak age windows, and BMI benchmarks for talent acquisition; and (2) Investigative Sports Media seeking historical national dominance trends, longitudinal gender equity metrics, and statistical proof of the host nation advantage.")

    # 2. Tech Stack Matrix
    add_h("2. Architecture & Applied Technology Stack", level=1)
    add_p("The solution was developed natively in Python and SQLite within Visual Studio Code, utilizing a modular toolchain:")
    tech_headers = ["Layer / Domain", "Library / Tool", "Practical Implementation Purpose"]
    tech_rows = [
        ["Database Engine", "SQLite3", "Serverless relational data storage, schema generation, and query execution engine."],
        ["Data Wrangling", "Pandas", "Ingested CSVs, verified schemas, and passed dataframes between SQL and visualizers."],
        ["Visual Analytics", "Matplotlib & Seaborn", "Rendered biometric clustering scatter plots and historical participation trends."],
        ["Reporting Engine", "python-docx", "Programmatically structured, styled, and generated executive Word deliverables."],
        ["Source Control", "Git & GitHub", "Version tracking, handling merge rebases, and public deployment of code assets."]
    ]
    create_table(tech_headers, tech_rows, [1.5, 1.8, 3.2])

    # 3. Data Hygiene & Entity Integrity
    add_h("3. Phase 1: Ingestion & Relational Hygiene Auditing", level=1)
    add_p("The raw dataset comprised two CSV files: athlete_events.csv (271,116 rows) and noc_regions.csv (230 rows). Ingestion and profiling revealed two primary structural challenges:")
    add_p("• Longitudinal Missing Data: Systematic auditing showed ~22–23% missing biometrics (Height: 60,171 nulls; Weight: 62,875 nulls). Nulls clustered predominantly in earlier Games (1896–1956). To preserve physiological integrity, naive mean imputation was rejected in favor of explicit biometric filtering (WHERE Height IS NOT NULL AND Weight IS NOT NULL).")
    add_p("• Broken Foreign Key Mismatches: Region mapping revealed orphan NOC codes (e.g., 'SGP' in event records vs. 'SIN' in regions). Solved using COALESCE(r.region, a.NOC) to guarantee zero dropped records during dimensional joins.")

    # 4. Core SQL Analytics & Methodologies
    add_h("4. Phase 2: Advanced SQL Engineering & Core Findings", level=1)
    
    add_h("A. Anthropometrics & BMI Profiling", level=2)
    add_p("Calculated dynamic Body Mass Index across athletic disciplines: BMI = Weight / (Height / 100)^2. Quantified high-BMI leverage sports (Weightlifting avg BMI: 27.6; Wrestling avg BMI: 25.7) versus high-efficiency power-to-weight sports (Gymnastics avg BMI: 21.2; Athletics avg BMI: 22.7).")

    add_h("B. Peak Competitive Age Windows", level=2)
    add_p("Evaluated gold medalist age distributions across sports. Identified highly front-loaded disciplines requiring early specialization (Swimming avg gold age: 21.7; Gymnastics: 22.1) versus precision and tactical sports enabling extended longevity (Equestrian avg gold age: 35.3; Shooting: 33.6).")

    add_h("C. Event Deduplication via CTEs (Resolving Roster Multiplicity)", level=2)
    add_p("Fact tables record medals on an athlete basis. In team sports (e.g., 1992 Men's Basketball), 12 gold medals were logged for 1 team victory. A naive COUNT(*) would artificially skew national dominance rankings. Implemented a Common Table Expression (CTE) using SELECT DISTINCT Year, Season, Sport, Event, Medal, NOC to evaluate true tournament wins.")

    add_h("D. Trailing Moving Window Functions (Host Nation Advantage)", level=2)
    add_p("Evaluated whether host nations experience an empirical performance boost. Constructed a 3-Games trailing moving average window function:")
    add_code_block(
        "ROUND(AVG(c.total_medals) OVER (\n"
        "    PARTITION BY c.country ORDER BY c.Year\n"
        "    ROWS BETWEEN 3 PRECEDING AND 1 PRECEDING\n"
        "), 1) AS trailing_3games_avg"
    )
    add_p("By excluding the current host year ('1 PRECEDING'), data leakage was prevented. The query empirically proved a +15 to +42 medal boost during host years (e.g., China +42.3 at Beijing 2008; Australia +30.7 at Sydney 2000).")

    # 5. Visual Storytelling & Automated Reporting
    add_h("5. Phase 3: Visual Analytics & Automated Reporting", level=1)
    add_p("• Visualization Exports: Built and saved two high-resolution standalone visual assets: 'biometric_clustering.png' (Height vs. Weight multi-sport scatter plot) and 'gender_growth.png' (1896–2016 female participation growth curve from 0% to 44.9%).")
    add_p("• Enterprise Document Automation: Executed generate_report.py to construct SportsStats_Final_Report.docx. Utilized python-docx to apply customized typography, alternating table shading, and programmatic figure embedding.")

    # 6. DevOps & Version Control
    add_h("6. Phase 4: DevOps, Git Troubleshooting & Public Deployment", level=1)
    add_p("The project was packaged and deployed to a public GitHub repository. Key version control operations included:")
    add_p("1. Identity Configuration: Resolved missing author errors via 'git config --global user.name' and 'user.email'.")
    add_p("2. Branch History Alignment: Overcame non-fast-forward push rejections caused by default remote commits using git push --force and git pull --rebase.")
    add_p("3. Local-to-Remote Synchronization: Cleanly rebased local Jupyter notebook modifications over newly authored Markdown documentation on GitHub without merge conflicts.")

    # 7. Interview & Portfolio Quick Reference
    add_h("7. Portfolio & Interview Talking Points", level=1)
    interview_headers = ["Technical Concept", "Interview Explanation / Project Context"]
    interview_rows = [
        ["Granularity & Multiplicity", "Explained how team rosters distort fact table aggregations, and demonstrated how CTEs with DISTINCT event keys establish accurate event-level metrics."],
        ["Analytical Windowing", "Articulated how ROWS BETWEEN 3 PRECEDING AND 1 PRECEDING models trailing baseline trends without leaking host year performance into the benchmark."],
        ["Data Cleansing Strategy", "Defended why mean imputation was rejected for historical biometrics due to chronological reporting bias, opting for strict conditional filtering."],
        ["Automated Delivery", "Demonstrated capability beyond analysis by writing Python automation pipelines (python-docx) to produce boardroom-ready executive deliverables."]
    ]
    create_table(interview_headers, interview_rows, [2.0, 4.5])

    doc.save("SportsStats_Project_Summary.docx")
    print("Complete Project Summary generated successfully: SportsStats_Project_Summary.docx")

if __name__ == "__main__":
    create_summary_document()